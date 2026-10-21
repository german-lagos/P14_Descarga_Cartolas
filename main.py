import os
import imaplib
import smtplib
import email

from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from email.header import decode_header
from dotenv import load_dotenv
from openpyxl import load_workbook


# ==========================================
# 1. CONFIGURACIÓN
# ==========================================

load_dotenv()

IMAP_SERVER = os.getenv("IMAP_SERVER")
IMAP_PORT = int(os.getenv("IMAP_PORT", "993"))

SMTP_SERVER = os.getenv("SMTP_SERVER")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))

EMAIL_USER = os.getenv("EMAIL_USER")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD")

EMAIL_REMITENTE = os.getenv("EMAIL_REMITENTE")
EMAIL_ASUNTO = os.getenv("EMAIL_ASUNTO")

EMAIL_DESTINATARIO = os.getenv("EMAIL_DESTINATARIO")

EMPRESA = os.getenv("EMPRESA")
FECHA_CARTOLA = os.getenv("FECHA_CARTOLA")

EXCEL_PATH = Path(os.getenv("EXCEL_PATH", "datos/cartolas.xlsx"))

DOWNLOAD_PATH = Path(
    os.getenv("DOWNLOAD_PATH", "descargas")
)

OUTPUT_PATH = Path(
    os.getenv("OUTPUT_PATH", "procesados")
)

LINEA_MODIFICAR = int(os.getenv("LINEA_MODIFICAR", "3"))
POSICION_INICIO = int(os.getenv("POSICION_INICIO", "16"))
LONGITUD_CAMPO = int(os.getenv("LONGITUD_CAMPO", "5"))

CODIFICACION = os.getenv(
    "CODIFICACION_ARCHIVO", "ascii"
)

MODO_PRUEBA = (
    os.getenv("MODO_PRUEBA", "true").lower() == "true"
)

DOWNLOAD_PATH.mkdir(parents=True, exist_ok=True)
OUTPUT_PATH.mkdir(parents=True, exist_ok=True)


# ==========================================
# 2. DESCARGAR ARCHIVO DAT
# ==========================================

def descargar_dat():

    print("\n[1] Buscando archivo DAT en el correo...")

    conexion = imaplib.IMAP4_SSL(
        IMAP_SERVER,
        IMAP_PORT
    )

    try:
        conexion.login(EMAIL_USER, EMAIL_PASSWORD)
        conexion.select("INBOX", readonly=True)

        estado, mensajes = conexion.search(
            None,
            "ALL"
        )

        if estado != "OK":
            raise RuntimeError("No se pudo consultar el buzón.")

        ids = mensajes[0].split()

        for mensaje_id in reversed(ids):

            estado, datos = conexion.fetch(
                mensaje_id,
                "(RFC822)"
            )

            if estado != "OK":
                continue

            mensaje = email.message_from_bytes(
                datos[0][1]
            )

            remitente = email.utils.parseaddr(
                mensaje.get("From", "")
            )[1]

            asunto, codificacion = decode_header(
                mensaje.get("Subject", "")
            )[0]

            if isinstance(asunto, bytes):
                asunto = asunto.decode(
                    codificacion or "utf-8",
                    errors="replace"
                )

            if remitente.lower() != EMAIL_REMITENTE.lower():
                continue

            if EMAIL_ASUNTO.lower() not in asunto.lower():
                continue

            for parte in mensaje.walk():

                nombre = parte.get_filename()

                if not nombre:
                    continue

                if not nombre.lower().endswith(".dat"):
                    continue

                nombre = os.path.basename(nombre)

                ruta = DOWNLOAD_PATH / nombre

                contenido = parte.get_payload(
                    decode=True
                )

                if contenido is None:
                    continue

                ruta.write_bytes(contenido)

                print(f"Archivo descargado: {ruta}")

                return ruta

        raise FileNotFoundError(
            "No se encontró un archivo DAT "
            "que cumpla los criterios."
        )

    finally:
        conexion.logout()


# ==========================================
# 3. CONSULTAR EXCEL
# ==========================================

def obtener_numero_cartola():

    print("\n[2] Consultando archivo Excel...")

    libro = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    try:
        hoja = libro.active

        fecha_buscada = datetime.strptime(
            FECHA_CARTOLA,
            "%Y-%m-%d"
        ).date()

        for fila in hoja.iter_rows(
            min_row=2,
            values_only=True
        ):

            empresa_excel = fila[0]
            fecha_excel = fila[1]
            numero_cartola = fila[2]

            if isinstance(fecha_excel, datetime):
                fecha_excel = fecha_excel.date()

            elif isinstance(fecha_excel, str):
                fecha_excel = datetime.strptime(
                    fecha_excel,
                    "%Y-%m-%d"
                ).date()

            if (
                str(empresa_excel).strip().lower()
                == EMPRESA.strip().lower()
                and fecha_excel == fecha_buscada
            ):

                if numero_cartola is None:
                    raise ValueError(
                        "El número de cartola está vacío."
                    )

                numero = str(numero_cartola)

                if not numero.isdigit():
                    raise ValueError(
                        "El número de cartola no es numérico."
                    )

                print(f"Empresa: {EMPRESA}")
                print(f"Fecha: {FECHA_CARTOLA}")
                print(f"Número de cartola: {numero}")

                return numero

        raise ValueError(
            "No se encontró la empresa y fecha en Excel."
        )

    finally:
        libro.close()


# ==========================================
# 4. MODIFICAR DAT Y GENERAR TXT
# ==========================================

def procesar_archivo(ruta_dat, numero_cartola):

    print("\n[3] Procesando archivo DAT...")

    # Se preservan los saltos de línea originales.
    with open(
        ruta_dat,
        "r",
        encoding=CODIFICACION,
        newline=""
    ) as archivo:

        contenido = archivo.readlines()

    indice_linea = LINEA_MODIFICAR - 1

    if indice_linea < 0 or indice_linea >= len(contenido):
        raise ValueError(
            "La línea indicada no existe en el archivo."
        )

    linea_original = contenido[indice_linea]

    if len(numero_cartola) > LONGITUD_CAMPO:
        raise ValueError(
            "El número de cartola supera "
            "la longitud permitida."
        )

    numero_formateado = numero_cartola.zfill(
        LONGITUD_CAMPO
    )

    inicio = POSICION_INICIO
    fin = inicio + LONGITUD_CAMPO

    if inicio < 0 or fin > len(linea_original.rstrip("\r\n")):
        raise ValueError(
            "La posición de reemplazo está "
            "fuera de la línea."
        )

    segmento_original = linea_original[inicio:fin]

    if not segmento_original.isdigit():
        raise ValueError(
            "El segmento original no es numérico."
        )

    linea_modificada = (
        linea_original[:inicio]
        + numero_formateado
        + linea_original[fin:]
    )

    contenido[indice_linea] = linea_modificada

    nombre_empresa = (
        EMPRESA.upper()
        .replace(" ", "_")
        .replace("/", "_")
    )

    fecha_nombre = FECHA_CARTOLA.replace("-", "")

    nombre_salida = (
        f"CARTOLA_{nombre_empresa}_{fecha_nombre}.txt"
    )

    ruta_salida = OUTPUT_PATH / nombre_salida

    with open(
        ruta_salida,
        "w",
        encoding=CODIFICACION,
        newline=""
    ) as archivo:

        archivo.writelines(contenido)

    print(f"Archivo generado: {ruta_salida}")
    print(f"Valor original: {segmento_original}")
    print(f"Valor nuevo: {numero_formateado}")

    return ruta_salida


# ==========================================
# 5. ENVIAR CORREO
# ==========================================

def enviar_correo(ruta_archivo):

    print("\n[4] Preparando correo electrónico...")

    mensaje = EmailMessage()

    mensaje["From"] = EMAIL_USER
    mensaje["To"] = EMAIL_DESTINATARIO

    mensaje["Subject"] = (
        f"Cartola procesada - {EMPRESA}"
    )

    mensaje.set_content(
        f"Se adjunta la cartola correspondiente "
        f"a {EMPRESA}, fecha {FECHA_CARTOLA}."
    )

    with open(ruta_archivo, "rb") as archivo:

        mensaje.add_attachment(
            archivo.read(),
            maintype="text",
            subtype="plain",
            filename=ruta_archivo.name
        )

    if MODO_PRUEBA:

        print("MODO PRUEBA ACTIVADO")
        print("El correo NO será enviado.")
        print(f"Destinatario: {EMAIL_DESTINATARIO}")
        print(f"Adjunto: {ruta_archivo.name}")

        return

    with smtplib.SMTP(
        SMTP_SERVER,
        SMTP_PORT
    ) as servidor:

        servidor.starttls()

        servidor.login(
            EMAIL_USER,
            EMAIL_PASSWORD
        )

        servidor.send_message(mensaje)

    print("Correo enviado correctamente.")


# ==========================================
# 6. FLUJO PRINCIPAL
# ==========================================

def main():

    print("=" * 55)
    print(" P14 - D2 DESCARGA DE CARTOLAS BANCO ESTADO")
    print("=" * 55)

    try:

        ruta_dat = descargar_dat()

        numero_cartola = obtener_numero_cartola()

        ruta_txt = procesar_archivo(
            ruta_dat,
            numero_cartola
        )

        enviar_correo(ruta_txt)

        print("\nProceso finalizado correctamente.")

    except Exception as error:

        print("\nERROR EN EL PROCESO:")
        print(error)


if __name__ == "__main__":
    main()