import os
from pathlib import Path
from datetime import datetime, date

from dotenv import load_dotenv
from openpyxl import load_workbook


# ==========================================
# 1. CONFIGURACIÓN
# ==========================================

BASE_DIR = Path(__file__).resolve().parent

load_dotenv(BASE_DIR / ".env")

EMPRESA = os.getenv("EMPRESA")
FECHA_CARTOLA = os.getenv("FECHA_CARTOLA")

EXCEL_PATH = BASE_DIR / os.getenv(
    "EXCEL_PATH", "datos/cartolas.xlsx"
)

DAT_PATH = BASE_DIR / os.getenv(
    "DAT_PATH", "descargas/CARTOLA_PRUEBA.DAT"
)

OUTPUT_PATH = BASE_DIR / os.getenv(
    "OUTPUT_PATH", "procesados"
)

LINEA_MODIFICAR = int(
    os.getenv("LINEA_MODIFICAR", "3")
)

POSICION_INICIO = int(
    os.getenv("POSICION_INICIO", "16")
)

LONGITUD_CAMPO = int(
    os.getenv("LONGITUD_CAMPO", "5")
)

CODIFICACION = os.getenv(
    "CODIFICACION_ARCHIVO", "ascii"
)


# ==========================================
# 2. VALIDAR CONFIGURACIÓN
# ==========================================

def validar_configuracion():

    print("\n[1] Validando configuración...")

    if not EMPRESA:
        raise ValueError("EMPRESA no está configurada.")

    if not FECHA_CARTOLA:
        raise ValueError("FECHA_CARTOLA no está configurada.")

    if not EXCEL_PATH.is_file():
        raise FileNotFoundError(
            f"No existe el Excel: {EXCEL_PATH}"
        )

    if not DAT_PATH.is_file():
        raise FileNotFoundError(
            f"No existe el DAT: {DAT_PATH}"
        )

    if LINEA_MODIFICAR < 1:
        raise ValueError("LINEA_MODIFICAR debe ser mayor que 0.")

    if POSICION_INICIO < 0:
        raise ValueError("POSICION_INICIO no puede ser negativa.")

    if LONGITUD_CAMPO < 1:
        raise ValueError("LONGITUD_CAMPO debe ser mayor que 0.")

    OUTPUT_PATH.mkdir(
        parents=True,
        exist_ok=True
    )

    print("Configuración validada correctamente.")


# ==========================================
# 3. CONSULTAR NÚMERO DE CARTOLA
# ==========================================

def obtener_numero_cartola():

    print("\n[2] Consultando archivo Excel...")

    fecha_buscada = datetime.strptime(
        FECHA_CARTOLA, "%Y-%m-%d"
    ).date()

    libro = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    try:
        hoja = libro.active

        encabezados = [
            str(celda.value).strip()
            if celda.value is not None else ""
            for celda in hoja[1]
        ]

        columnas = {
            nombre: indice
            for indice, nombre in enumerate(encabezados)
        }

        requeridas = [
            "Empresa",
            "Fecha",
            "Numero_Cartola"
        ]

        for columna in requeridas:
            if columna not in columnas:
                raise ValueError(
                    f"Falta la columna '{columna}' en Excel."
                )

        coincidencias = []

        for fila in hoja.iter_rows(
            min_row=2,
            values_only=True
        ):

            empresa_excel = fila[columnas["Empresa"]]
            fecha_excel = fila[columnas["Fecha"]]
            numero_excel = fila[columnas["Numero_Cartola"]]

            if empresa_excel is None or fecha_excel is None:
                continue

            if isinstance(fecha_excel, datetime):
                fecha_excel = fecha_excel.date()

            elif isinstance(fecha_excel, str):
                fecha_excel = datetime.strptime(
                    fecha_excel.strip(),
                    "%Y-%m-%d"
                ).date()

            elif not isinstance(fecha_excel, date):
                raise ValueError(
                    f"Formato de fecha no válido: {fecha_excel}"
                )

            if (
                str(empresa_excel).strip().casefold()
                == EMPRESA.strip().casefold()
                and fecha_excel == fecha_buscada
            ):

                if numero_excel is None:
                    raise ValueError(
                        "El número de cartola está vacío."
                    )

                # Excel puede entregar valores numéricos
                # como enteros o decimales.
                if isinstance(numero_excel, int):
                    numero = str(numero_excel)

                elif isinstance(numero_excel, float):
                    if not numero_excel.is_integer():
                        raise ValueError(
                            "El número de cartola contiene decimales."
                        )
                    numero = str(int(numero_excel))

                else:
                    numero = str(numero_excel).strip()

                if not numero.isascii() or not numero.isdigit():
                    raise ValueError(
                        "El número de cartola debe contener "
                        "únicamente dígitos del 0 al 9."
                    )

                coincidencias.append(numero)

        if not coincidencias:
            raise ValueError(
                "No se encontró una cartola para "
                f"{EMPRESA} en la fecha {FECHA_CARTOLA}."
            )

        if len(coincidencias) > 1:
            raise ValueError(
                "Se encontraron múltiples registros "
                "para la misma empresa y fecha."
            )

        numero = coincidencias[0]

        print(f"Empresa: {EMPRESA}")
        print(f"Fecha: {FECHA_CARTOLA}")
        print(f"Número de cartola: {numero}")

        return numero

    finally:
        libro.close()


# ==========================================
# 4. MODIFICAR DAT Y GENERAR TXT
# ==========================================

def procesar_dat(numero_cartola):

    print("\n[3] Procesando archivo DAT...")

    # newline="" evita transformar los saltos
    # de línea originales del archivo.
    with open(
        DAT_PATH,
        "r",
        encoding=CODIFICACION,
        newline=""
    ) as archivo:

        lineas = archivo.readlines()

    indice = LINEA_MODIFICAR - 1

    if indice >= len(lineas):
        raise ValueError(
            f"El archivo tiene solamente {len(lineas)} líneas."
        )

    linea_original = lineas[indice]

    # Separar contenido y salto de línea.
    contenido_linea = linea_original.rstrip("\r\n")
    salto_linea = linea_original[len(contenido_linea):]

    inicio = POSICION_INICIO
    fin = inicio + LONGITUD_CAMPO

    if fin > len(contenido_linea):
        raise ValueError(
            "El campo indicado supera la longitud de la línea."
        )

    if len(numero_cartola) > LONGITUD_CAMPO:
        raise ValueError(
            "El número de cartola supera la longitud del campo."
        )

    numero_formateado = numero_cartola.zfill(
        LONGITUD_CAMPO
    )

    valor_original = contenido_linea[inicio:fin]

    if not valor_original.isascii() or not valor_original.isdigit():
        raise ValueError(
            "El campo original no contiene únicamente números."
        )

    contenido_modificado = (
        contenido_linea[:inicio]
        + numero_formateado
        + contenido_linea[fin:]
    )

    lineas[indice] = contenido_modificado + salto_linea

    # Verificar que no cambió la longitud.
    if len(lineas[indice]) != len(linea_original):
        raise ValueError(
            "La modificación alteró la longitud del registro."
        )

    nombre_empresa = "".join(
        caracter if caracter.isalnum() else "_"
        for caracter in EMPRESA.upper()
    )

    fecha_nombre = FECHA_CARTOLA.replace("-", "")

    nombre_salida = (
        f"CARTOLA_{nombre_empresa}_{fecha_nombre}.txt"
    )

    ruta_salida = OUTPUT_PATH / nombre_salida

    # Evitar sobrescribir resultados anteriores.
    if ruta_salida.exists():
        raise FileExistsError(
            f"El archivo ya existe: {ruta_salida}"
        )

    with open(
        ruta_salida,
        "w",
        encoding=CODIFICACION,
        newline=""
    ) as archivo:

        archivo.writelines(lineas)

    print(f"Archivo original: {DAT_PATH.name}")
    print(f"Valor anterior: {valor_original}")
    print(f"Valor nuevo: {numero_formateado}")
    print(f"Archivo generado: {ruta_salida}")

    return ruta_salida


# ==========================================
# 5. EJECUCIÓN PRINCIPAL
# ==========================================

def main():

    print("=" * 55)
    print(" P14 - PRUEBA DE PROCESAMIENTO DAT A TXT")
    print("=" * 55)

    try:

        validar_configuracion()

        numero_cartola = obtener_numero_cartola()

        ruta_txt = procesar_dat(numero_cartola)

        print("\n[4] Proceso completado correctamente.")
        print(f"Resultado: {ruta_txt.resolve()}")

    except Exception as error:

        print("\nERROR EN EL PROCESO:")
        print(f"{type(error).__name__}: {error}")


if __name__ == "__main__":
    main()